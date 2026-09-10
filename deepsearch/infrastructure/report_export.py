"""把规范 Markdown 报告按需导出为常用文件格式。"""

from __future__ import annotations

import html
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit


REPORT_EXPORT_FORMATS = (
    ("markdown", "Markdown"),
    ("docx", "Word"),
    ("pdf", "PDF"),
    ("text", "TXT"),
    ("json", "JSON"),
    ("html", "HTML"),
)


class ReportExportError(RuntimeError):
    """报告无法转换为请求格式。"""


@dataclass(frozen=True, slots=True)
class ExportedReport:
    """可直接交给下载响应的文件内容与元数据。"""

    data: bytes
    file_name: str
    mime_type: str
    format_key: str
    format_label: str


@dataclass(frozen=True, slots=True)
class _Block:
    kind: str
    text: str = ""
    level: int = 0
    items: tuple[str, ...] = ()
    rows: tuple[tuple[str, ...], ...] = ()


_FORMAT_METADATA = {
    "markdown": ("Markdown", ".md", "text/markdown"),
    "docx": ("Word", ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": ("PDF", ".pdf", "application/pdf"),
    "text": ("TXT", ".txt", "text/plain"),
    "json": ("JSON", ".json", "application/json"),
    "html": ("HTML", ".html", "text/html"),
}
_TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?$")
_ORDERED_ITEM = re.compile(r"^\s*\d+[.)]\s+(.+)$")
_UNORDERED_ITEM = re.compile(r"^\s*[-*+]\s+(.+)$")
_HEADING = re.compile(r"^(#{1,6})\s+(.+)$")
_MARKDOWN_LINK = re.compile(r"\[([^]]+)]\(([^)]+)\)")


def export_report(markdown: str, title: str, output_format: str) -> ExportedReport:
    """从同一份 Markdown 生成一个下载文件，不写入资料库。"""

    normalized = output_format.strip().lower()
    if normalized not in _FORMAT_METADATA:
        raise ValueError(f"不支持的导出格式：{output_format}")
    label, extension, mime_type = _FORMAT_METADATA[normalized]
    clean_title = title.strip() or _document_title(markdown) or "报告"
    blocks = _parse_markdown(markdown)

    if normalized == "markdown":
        data = markdown.encode("utf-8")
    elif normalized == "text":
        data = _to_text(blocks).encode("utf-8")
    elif normalized == "json":
        data = json.dumps(
            {
                "schema_version": "1.0",
                "question": clean_title,
                "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
                "content_markdown": markdown,
            },
            ensure_ascii=False,
            indent=2,
        ).encode("utf-8")
    elif normalized == "html":
        data = _to_html(blocks, clean_title).encode("utf-8")
    elif normalized == "docx":
        data = _to_docx(blocks, clean_title)
    else:
        data = _to_pdf(blocks, clean_title)

    return ExportedReport(
        data=data,
        file_name=f"{_safe_file_stem(clean_title)}{extension}",
        mime_type=mime_type,
        format_key=normalized,
        format_label=label,
    )


def _document_title(markdown: str) -> str:
    return next(
        (match.group(2).strip() for line in markdown.splitlines() if (match := _HEADING.match(line.strip()))),
        "",
    )


def _safe_file_stem(title: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", title)
    value = re.sub(r"\s+", " ", value).strip(" .")
    return (value[:96].rstrip(" .") or "报告")


def _parse_markdown(markdown: str) -> list[_Block]:
    """解析导出所需的受控 Markdown 子集，保持所有格式语义一致。"""

    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[_Block] = []
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped.removeprefix("```").strip()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1 if index < len(lines) else 0
            blocks.append(_Block("code", "\n".join(code_lines), items=(language,)))
            continue

        heading = _HEADING.match(stripped)
        if heading:
            blocks.append(_Block("heading", heading.group(2).strip(), len(heading.group(1))))
            index += 1
            continue

        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and _TABLE_SEPARATOR.match(lines[index + 1].strip())
        ):
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = tuple(_split_table_row(value) for value in table_lines)
            width = max((len(row) for row in rows), default=0)
            normalized_rows = tuple(row + ("",) * (width - len(row)) for row in rows)
            blocks.append(_Block("table", rows=normalized_rows))
            continue

        unordered = _UNORDERED_ITEM.match(line)
        if unordered:
            items = []
            while index < len(lines) and (match := _UNORDERED_ITEM.match(lines[index])):
                items.append(match.group(1).strip())
                index += 1
            blocks.append(_Block("unordered", items=tuple(items)))
            continue

        ordered = _ORDERED_ITEM.match(line)
        if ordered:
            items = []
            while index < len(lines) and (match := _ORDERED_ITEM.match(lines[index])):
                items.append(match.group(1).strip())
                index += 1
            blocks.append(_Block("ordered", items=tuple(items)))
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().removeprefix(">").strip())
                index += 1
            blocks.append(_Block("quote", "\n".join(quote_lines)))
            continue

        if re.fullmatch(r"[-*_]{3,}", stripped):
            blocks.append(_Block("rule"))
            index += 1
            continue

        paragraph = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or _starts_new_block(lines, index):
                break
            paragraph.append(candidate)
            index += 1
        blocks.append(_Block("paragraph", " ".join(paragraph)))
    return blocks


def _starts_new_block(lines: list[str], index: int) -> bool:
    value = lines[index].strip()
    return bool(
        value.startswith(("```", ">"))
        or _HEADING.match(value)
        or _UNORDERED_ITEM.match(value)
        or _ORDERED_ITEM.match(value)
        or re.fullmatch(r"[-*_]{3,}", value)
        or (
            value.startswith("|")
            and index + 1 < len(lines)
            and _TABLE_SEPARATOR.match(lines[index + 1].strip())
        )
    )


def _split_table_row(line: str) -> tuple[str, ...]:
    return tuple(cell.strip() for cell in line.strip().strip("|").split("|"))


def _plain_inline(value: str) -> str:
    value = re.sub(r"!\[([^]]*)]\([^)]+\)", r"\1", value)
    value = _MARKDOWN_LINK.sub(lambda match: f"{match.group(1)}（{match.group(2)}）", value)
    value = re.sub(r"(\*\*|__)(.+?)\1", r"\2", value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", value)
    value = re.sub(r"(?<!_)_([^_]+)_(?!_)", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    return value


def _pdf_safe_text(value: str) -> str:
    replacements = {
        "✅": "通过",
        "⚠️": "注意",
        "⚠": "注意",
        "❌": "冲突",
        "🧪": "模拟",
    }
    for marker, replacement in replacements.items():
        value = value.replace(marker, replacement)
    return _plain_inline(value)


def _to_text(blocks: list[_Block]) -> str:
    lines: list[str] = []
    for block in blocks:
        if block.kind == "heading":
            lines.extend([_plain_inline(block.text), ""])
        elif block.kind in {"paragraph", "quote"}:
            lines.extend([_plain_inline(block.text), ""])
        elif block.kind == "unordered":
            lines.extend(f"- {_plain_inline(item)}" for item in block.items)
            lines.append("")
        elif block.kind == "ordered":
            lines.extend(f"{index}. {_plain_inline(item)}" for index, item in enumerate(block.items, 1))
            lines.append("")
        elif block.kind == "code":
            lines.extend([block.text, ""])
        elif block.kind == "table":
            lines.extend("\t".join(_plain_inline(cell) for cell in row) for row in block.rows)
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def _inline_html(value: str) -> str:
    """渲染安全的行内标记；报告链接只允许普通 Web 协议。"""

    value = re.sub(r"!\[([^]]*)]\([^)]+\)", r"\1", value)
    output: list[str] = []
    cursor = 0
    for match in _MARKDOWN_LINK.finditer(value):
        output.append(html.escape(value[cursor:match.start()]))
        label = html.escape(_plain_inline(match.group(1)))
        raw_url = match.group(2).strip()
        if urlsplit(raw_url).scheme.lower() in {"http", "https"}:
            url = html.escape(raw_url, quote=True)
            output.append(f'<a href="{url}" rel="noopener noreferrer">{label}</a>')
        else:
            # 导出文件可能被用户直接在浏览器中打开；未知协议保留可读
            # 标签但不生成可点击链接，避免 javascript:/data: 被执行。
            output.append(label)
        cursor = match.end()
    output.append(html.escape(value[cursor:]))
    rendered = "".join(output)
    rendered = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", rendered)
    rendered = re.sub(r"__(.+?)__", r"<strong>\1</strong>", rendered)
    rendered = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", rendered)
    rendered = re.sub(r"(?<!_)_([^_]+)_(?!_)", r"<em>\1</em>", rendered)
    rendered = re.sub(r"`([^`]+)`", r"<code>\1</code>", rendered)
    return rendered


def _to_html(blocks: list[_Block], title: str) -> str:
    """生成可离线阅读的自包含 HTML，不引用第三方脚本或样式。"""

    body: list[str] = []
    for block in blocks:
        if block.kind == "heading":
            level = min(max(block.level, 1), 6)
            body.append(f"<h{level}>{_inline_html(block.text)}</h{level}>")
        elif block.kind == "paragraph":
            body.append(f"<p>{_inline_html(block.text)}</p>")
        elif block.kind == "quote":
            body.append(f"<blockquote>{'<br>'.join(_inline_html(line) for line in block.text.splitlines())}</blockquote>")
        elif block.kind in {"unordered", "ordered"}:
            tag = "ul" if block.kind == "unordered" else "ol"
            items = "".join(f"<li>{_inline_html(item)}</li>" for item in block.items)
            body.append(f"<{tag}>{items}</{tag}>")
        elif block.kind == "code":
            body.append(f"<pre><code>{html.escape(block.text)}</code></pre>")
        elif block.kind == "rule":
            body.append("<hr>")
        elif block.kind == "table" and block.rows:
            header_cells = "".join(f"<th>{_inline_html(cell)}</th>" for cell in block.rows[0])
            rows = [f"<thead><tr>{header_cells}</tr></thead>"]
            rows.extend(
                "<tr>" + "".join(f"<td>{_inline_html(cell)}</td>" for cell in row) + "</tr>"
                for row in block.rows[1:]
            )
            body.append(f"<div class=\"table-wrap\"><table>{''.join(rows)}</table></div>")
    return """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
body {{ margin: 0 auto; max-width: 900px; padding: 48px 32px 72px; color: #172033;
       font: 16px/1.75 system-ui, "Microsoft YaHei", "PingFang SC", sans-serif; }}
h1 {{ margin: 0 0 28px; font-size: 2rem; line-height: 1.35; }}
h2 {{ margin: 34px 0 12px; font-size: 1.42rem; }}
h3 {{ margin: 26px 0 10px; font-size: 1.16rem; }}
p, li {{ margin: 8px 0; }}
a {{ color: #087f73; text-underline-offset: 2px; }}
blockquote {{ margin: 18px 0; padding: 12px 18px; border-left: 3px solid #98a2b3; color: #475467; }}
pre {{ overflow-x: auto; padding: 16px; border-radius: 10px; background: #f4f6f8; }}
code {{ font-family: ui-monospace, "Cascadia Code", monospace; }}
.table-wrap {{ overflow-x: auto; margin: 20px 0; }}
table {{ width: 100%; border-collapse: collapse; font-size: .94rem; }}
th, td {{ padding: 10px 12px; border: 1px solid #d9d9d9; text-align: left; vertical-align: middle; }}
th {{ background: #243b53; color: white; }}
tbody tr:nth-child(even) {{ background: #f7f9fb; }}
@media (max-width: 640px) {{ body {{ padding: 24px 18px 48px; }} }}
</style>
</head>
<body>{body}</body>
</html>
""".format(title=html.escape(title), body="\n".join(body))


def _to_docx(blocks: list[_Block], title: str) -> bytes:
    """按需生成 Word；资料库仍以原始 Markdown 作为单一事实来源。"""

    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise ReportExportError("Word 导出需要安装 python-docx。") from exc

    document = Document()
    document.core_properties.title = title
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.75)
    section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.82)
    section.right_margin = Inches(.82)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.styles["Normal"].paragraph_format.line_spacing = 1.35
    document.styles["Title"].font.size = Pt(23)
    document.styles["Title"].font.bold = True
    document.styles["Title"].paragraph_format.space_after = Pt(16)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(13)
    document.styles["Heading 3"].font.size = Pt(11.5)

    title_written = False
    for block in blocks:
        if block.kind == "heading":
            text = _plain_inline(block.text)
            if block.level == 1 and not title_written:
                paragraph = document.add_paragraph(text, style="Title")
                title_written = True
            else:
                level = min(max(block.level - 1, 1), 3)
                paragraph = document.add_heading(text, level=level)
            paragraph.paragraph_format.keep_with_next = True
        elif block.kind == "paragraph":
            document.add_paragraph(_plain_inline(block.text))
        elif block.kind == "quote":
            for line in block.text.splitlines():
                paragraph = document.add_paragraph(_plain_inline(line))
                paragraph.paragraph_format.left_indent = Inches(.22)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.style = document.styles["Normal"]
        elif block.kind in {"unordered", "ordered"}:
            style = "List Bullet" if block.kind == "unordered" else "List Number"
            for item in block.items:
                document.add_paragraph(_plain_inline(item), style=style)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(.18)
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif block.kind == "rule":
            document.add_paragraph()
        elif block.kind == "table" and block.rows:
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for row_index, values in enumerate(block.rows):
                for column_index, value in enumerate(values):
                    cell = table.cell(row_index, column_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = _plain_inline(value)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        paragraph.paragraph_format.space_before = Pt(3)
                        for run in paragraph.runs:
                            run.font.name = "Microsoft YaHei"
                            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                            run.font.size = Pt(8.5)
                            if row_index == 0:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "243B53")
                        cell._tc.get_or_add_tcPr().append(shading)
            document.add_paragraph().paragraph_format.space_after = Pt(0)

    if not title_written:
        document.paragraphs[0].insert_paragraph_before(title, style="Title") if document.paragraphs else document.add_paragraph(title, style="Title")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _to_pdf(blocks: list[_Block], title: str) -> bytes:
    """按需生成 PDF，并优先选择可嵌入的跨平台中文字体。"""

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise ReportExportError("PDF 导出需要安装 reportlab。") from exc

    font_name = _register_pdf_font(pdfmetrics, TTFont, UnicodeCIDFont)

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ChineseBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=17,
        textColor=colors.HexColor("#172033"),
        spaceAfter=7,
        splitLongWords=True,
        wordWrap="CJK",
    )
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=body,
        fontSize=22,
        leading=30,
        spaceAfter=18,
        keepWithNext=True,
    )
    heading_styles = {
        1: ParagraphStyle("ChineseH1", parent=body, fontSize=16, leading=22, spaceBefore=15, spaceAfter=8, keepWithNext=True),
        2: ParagraphStyle("ChineseH2", parent=body, fontSize=13, leading=19, spaceBefore=12, spaceAfter=6, keepWithNext=True),
        3: ParagraphStyle("ChineseH3", parent=body, fontSize=11.5, leading=17, spaceBefore=9, spaceAfter=5, keepWithNext=True),
    }
    quote_style = ParagraphStyle(
        "ChineseQuote",
        parent=body,
        leftIndent=12,
        textColor=colors.HexColor("#475467"),
        borderColor=colors.HexColor("#98A2B3"),
        borderWidth=0,
        borderPadding=(0, 0, 0, 8),
    )
    code_style = ParagraphStyle(
        "ChineseCode",
        parent=body,
        fontSize=8.5,
        leading=12,
        leftIndent=10,
        backColor=colors.HexColor("#F4F6F8"),
        borderPadding=8,
    )
    list_style = ParagraphStyle(
        "ChineseList",
        parent=body,
        leftIndent=16,
        firstLineIndent=-12,
        spaceAfter=4,
    )
    table_body_style = ParagraphStyle(
        "ChineseTableBody",
        parent=body,
        fontSize=8.5,
        leading=12,
        spaceAfter=0,
    )
    table_header_style = ParagraphStyle(
        "ChineseTableHeader",
        parent=table_body_style,
        textColor=colors.white,
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        title=title,
        leftMargin=.78 * inch,
        rightMargin=.78 * inch,
        topMargin=.72 * inch,
        bottomMargin=.72 * inch,
    )
    story = []
    title_written = False
    for block in blocks:
        if block.kind == "heading":
            text = html.escape(_pdf_safe_text(block.text))
            if block.level == 1 and not title_written:
                story.append(Paragraph(text, title_style))
                title_written = True
            else:
                story.append(Paragraph(text, heading_styles[min(max(block.level - 1, 1), 3)]))
        elif block.kind == "paragraph":
            story.append(Paragraph(html.escape(_pdf_safe_text(block.text)), body))
        elif block.kind == "quote":
            story.append(Paragraph("<br/>".join(html.escape(_pdf_safe_text(line)) for line in block.text.splitlines()), quote_style))
        elif block.kind in {"unordered", "ordered"}:
            for item_index, item in enumerate(block.items, 1):
                # The middle dot is covered by the bundled/system CJK fonts more
                # consistently than the Unicode bullet on Windows PDF readers.
                marker = "·" if block.kind == "unordered" else f"{item_index}."
                story.append(Paragraph(
                    f"{marker} {html.escape(_pdf_safe_text(item))}",
                    list_style,
                ))
            story.append(Spacer(1, 5))
        elif block.kind == "code":
            story.append(Paragraph(html.escape(_pdf_safe_text(block.text)).replace("\n", "<br/>"), code_style))
        elif block.kind == "rule":
            story.append(Spacer(1, 8))
        elif block.kind == "table" and block.rows:
            table_data = [
                [
                    Paragraph(
                        html.escape(_pdf_safe_text(cell)),
                        table_header_style if row_index == 0 else table_body_style,
                    )
                    for cell in row
                ]
                for row_index, row in enumerate(block.rows)
            ]
            column_count = len(block.rows[0])
            available_width = 6.94 * inch
            weights = []
            for column_index in range(column_count):
                longest = max(len(_pdf_safe_text(row[column_index])) for row in block.rows)
                weights.append(min(max(longest, 8), 36))
            weight_total = sum(weights) or column_count
            widths = [available_width * weight / weight_total for weight in weights]
            table = Table(table_data, colWidths=widths, repeatRows=1, hAlign="CENTER")
            table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B53")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FB")]),
            ]))
            story.extend([table, Spacer(1, 10)])

    if not title_written:
        story.insert(0, Paragraph(html.escape(_pdf_safe_text(title)), title_style))
    document.build(story)
    return output.getvalue()


def _register_pdf_font(pdfmetrics, tt_font_class, cid_font_class) -> str:
    """优先嵌入跨阅读器稳定的中文字体，找不到时回退到标准 CID 字体。"""

    project_font = Path(__file__).resolve().parents[2] / "assets" / "fonts" / "NotoSansSC-Regular.ttf"
    windows_root = Path(os.environ.get("WINDIR", "C:/Windows"))
    candidates = (
        project_font,
        windows_root / "Fonts" / "Deng.ttf",
        windows_root / "Fonts" / "simhei.ttf",
        windows_root / "Fonts" / "simsunb.ttf",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.otf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansSC-Regular.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
    )
    for candidate in candidates:
        if not candidate.is_file():
            continue
        try:
            pdfmetrics.registerFont(tt_font_class("DeepSearchCJK", str(candidate)))
        except Exception:
            continue
        return "DeepSearchCJK"

    fallback = "STSong-Light"
    try:
        pdfmetrics.getFont(fallback)
    except KeyError:
        pdfmetrics.registerFont(cid_font_class(fallback))
    return fallback
