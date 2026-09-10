import json
import unittest
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from deepsearch.infrastructure.report_export import REPORT_EXPORT_FORMATS, export_report


SAMPLE_REPORT = """# 中文测试报告

> 生成时间：2026-09-10 10:00

## 结论

这是包含 **重点结论** 和 [OpenAI](https://openai.com/) 链接的正文。

- 第一项
- 第二项

### 对比表

| 项目 | 结果 |
|---|---|
| 中文内容 | 正常 |
| English | OpenAI |
"""


class ReportExportTests(unittest.TestCase):
    def test_text_json_html_and_markdown_exports_keep_report_content(self):
        markdown = export_report(SAMPLE_REPORT, "中文测试报告", "markdown")
        text = export_report(SAMPLE_REPORT, "中文测试报告", "text")
        structured = export_report(SAMPLE_REPORT, "中文测试报告", "json")
        webpage = export_report(SAMPLE_REPORT, "中文测试报告", "html")

        self.assertEqual(markdown.data.decode("utf-8"), SAMPLE_REPORT)
        self.assertTrue(markdown.file_name.endswith(".md"))
        self.assertIn("重点结论", text.data.decode("utf-8"))
        self.assertNotIn("**", text.data.decode("utf-8"))
        payload = json.loads(structured.data.decode("utf-8"))
        self.assertEqual(payload["question"], "中文测试报告")
        self.assertEqual(payload["content_markdown"], SAMPLE_REPORT)
        html_text = webpage.data.decode("utf-8")
        self.assertIn('<html lang="zh-CN">', html_text)
        self.assertIn(
            '<a href="https://openai.com/" rel="noopener noreferrer">OpenAI</a>',
            html_text,
        )
        self.assertIn("<table>", html_text)

    def test_html_export_does_not_create_executable_non_web_links(self):
        exported = export_report("# 报告\n\n[打开](javascript:alert(1))", "报告", "html")

        html_text = exported.data.decode("utf-8")
        self.assertNotIn("javascript:", html_text)
        self.assertNotIn("<a href=", html_text)
        self.assertIn("打开", html_text)

    def test_word_export_is_valid_docx_with_headings_lists_and_table(self):
        exported = export_report(SAMPLE_REPORT, "中文测试报告", "docx")

        self.assertEqual(exported.data[:2], b"PK")
        self.assertEqual(exported.mime_type, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        with ZipFile(BytesIO(exported.data)) as archive:
            self.assertIn("word/document.xml", archive.namelist())
        document = Document(BytesIO(exported.data))
        paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("中文测试报告", paragraph_text)
        self.assertIn("第一项", paragraph_text)
        self.assertEqual(document.tables[0].cell(1, 0).text, "中文内容")
        self.assertEqual(document.tables[0].cell(1, 1).text, "正常")

    def test_pdf_export_is_readable_and_contains_expected_ascii_text(self):
        exported = export_report(SAMPLE_REPORT, "中文测试报告", "pdf")

        self.assertTrue(exported.data.startswith(b"%PDF-"))
        reader = PdfReader(BytesIO(exported.data))
        self.assertGreaterEqual(len(reader.pages), 1)
        extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
        self.assertIn("OpenAI", extracted)

    def test_every_advertised_format_exports_with_a_safe_file_name(self):
        for format_key, _ in REPORT_EXPORT_FORMATS:
            with self.subTest(format_key=format_key):
                exported = export_report(SAMPLE_REPORT, '测试/报告:*?"', format_key)
                self.assertTrue(exported.data)
                self.assertNotRegex(exported.file_name, r'[/:*?"]')


if __name__ == "__main__":
    unittest.main()
